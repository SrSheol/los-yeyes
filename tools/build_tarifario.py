#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Karún Travel Group — build assets/tarifario.json + assets/tarifario.data.js
from the client's official rate workbook (.xlsm) and cross-check every price
against the published PDFs.

Usage (from repo root):
  python3 tools/build_tarifario.py \
      --xlsm  "/path/Tarifario_regular_Karun_Travel__Definitivo.xlsm" \
      --pdf-dir "/path/with/the/pdfs"          # optional, needs `pdftotext`

Only the four visible "DESDE ..." sheets are read (the older dated/hidden
sheets are ignored on purpose). Nothing is invented: a hotel only gets the
vehicle columns that its sheet publishes.
"""
import argparse, json, os, re, subprocess, sys, unicodedata, glob, datetime
import openpyxl

SHEETS = {
    # sheet name (exact, incl. trailing spaces) -> (origin, classes in columns C, D)
    "DESDE PUNTA CANA SUB CADILLAC ": ("PUJ", ("suburban", "cadillac")),
    "DESDE AEROP PUNTA CANA MINIVAN": ("PUJ", ("starex", "hiace")),
    "DESDE AEROP SDQ ": ("SDQ", ("starex", "hiace")),
    "DESDE AEROP SDQ  SUB CAD": ("SDQ", ("suburban", "cadillac")),
}

# Obvious spelling slips in the client sheets (same hotel, spelled two ways).
# Prices are never touched. Every fix is listed in meta.normalizations.
NAME_FIXES = {
    "OCCIENTAL CARIBE": "OCCIDENTAL CARIBE",
    "COURYARD BY MARRIOTT SDQ": "COURTYARD BY MARRIOTT SDQ",
}

ACRONYMS = {"AC", "W", "TRS", "PPC", "PPR", "SDQ", "PUJ", "JW", "HM", "P.C.", "P."}
LOWER = {"AT", "BY", "DE", "DEL", "LA", "EL", "&"}

def zlabel(zraw):
    return ZONE_LABELS.get(strip_accents(zraw), zraw.title())


ZONE_LABELS = {
    "UVERO ALTO": "Uvero Alto", "UVERO ALTO AREA": "Uvero Alto",
    "MACAO AREA": "Macao", "ARENA GORDA": "Arena Gorda", "BAVARO AREA": "Bávaro",
    "CABEZA DE TORO": "Cabeza de Toro", "CABEZA DE TORO ZONE (BAVARO ZONE)": "Cabeza de Toro",
    "CAP CANA": "Cap Cana", "CAP CANA ZONE (BAVARO ZONE)": "Cap Cana",
    "PUNTA CANA": "Punta Cana", "PUNTA CANA ZONE (BAVARO ZONE)": "Punta Cana",
    "PLAYA NUEVA ROMANA ZONE": "Playa Nueva Romana", "LA ROMANA ZONE": "La Romana",
    "BAYAHIBE ZONE": "Bayahibe", "JUAN DOLIO ZONE": "Juan Dolio",
    "BOCA CHICA ZONE": "Boca Chica", "SANTO DOMINGO ZONE": "Santo Domingo",
    "BAVARO ZONE I (BAVARO ZONE)": "Bávaro · Zona I",
    "BAVARO ZONE II (BAVARO ZONE)": "Bávaro · Zona II",
    "BAVARO ZONE III (BAVARO ZONE)": "Bávaro · Zona III",
}

VEHICLES = [
    {"id": "starex", "label": "Hyundai Starex", "paxMin": 1, "paxMax": 6,
     "sheetLabels": ["STAREX HYUNDAI", "STAREX"]},
    {"id": "hiace", "label": "Toyota Hiace · Techo alto", "labelEn": "Toyota Hiace · High roof",
     "paxMin": 7, "paxMax": 10, "sheetLabels": ["HIACE TOYOTA", "TECHO ALTO"]},
    {"id": "suburban", "label": "Chevrolet Suburban", "paxMin": 1, "paxMax": 4,
     "sheetLabels": ["SUBURBAN"]},
    {"id": "cadillac", "label": "Cadillac Escalade", "paxMin": 1, "paxMax": 4,
     "sheetLabels": ["CADILLAC"]},
]

ORIGINS = {
    "PUJ": {"id": "PUJ", "name": "Aeropuerto de Punta Cana (PUJ)", "nameEn": "Punta Cana Airport (PUJ)"},
    "SDQ": {"id": "SDQ", "name": "Aeropuerto de Santo Domingo (SDQ)", "nameEn": "Santo Domingo Airport (SDQ)"},
}


def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def clean_name(raw):
    s = re.sub(r"\s+", " ", str(raw)).strip()
    s = re.sub(r"^\d+\s+(?=[A-Za-zÁÉÍÓÚ])", "", s)          # SDQ sheet list numbering "1 EMBASSY..."
    s = s.upper()
    final = NAME_FIXES.get(s, s)
    return final, (final != re.sub(r"\s+", " ", str(raw)).strip().upper())


def display_name(up):
    out = []
    for w in up.split(" "):
        if w in ACRONYMS:
            out.append(w)
        elif w in LOWER:
            out.append(w.lower())
        else:
            out.append("-".join(p[:1] + p[1:].lower() for p in w.split("-")))
    s = " ".join(out)
    return s[:1].upper() + s[1:]


def slug(s):
    s = strip_accents(s).lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def read_sheet(ws, classes):
    """-> ordered list of (zone_raw, [(name_up, {cls: price}, {cls: extra})])"""
    zones, cur = [], None
    for r in range(5, ws.max_row + 1):
        name, a, b, e = (ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value)
        if name is None or str(name).strip() == "":
            continue
        if a is None and b is None:                          # zone header row
            cur = (re.sub(r"\s+", " ", str(name)).strip().upper(), [])
            zones.append(cur)
            continue
        assert cur is not None, f"row {r}: hotel before any zone"
        nm, fixed = clean_name(name)
        cur[1].append((nm, {classes[0]: float(a), classes[1]: float(b)},
                       {classes[0]: float(e), classes[1]: float(e)}, re.sub(r"\s+", " ", str(name)).strip() if fixed else None))
    return zones


def build(xlsm):
    wb = openpyxl.load_workbook(xlsm, data_only=True)
    origins = {}
    notes = {"normalizations": [], "duplicatesDropped": [], "flags": []}
    for sheet, (oid, classes) in SHEETS.items():
        zones = read_sheet(wb[sheet], classes)
        o = origins.setdefault(oid, {"zones": {}, "order": []})
        for zraw, rows in zones:
            z = o["zones"].get(zraw)
            if z is None:
                z = o["zones"][zraw] = {"hotels": {}, "order": []}
                o["order"].append(zraw)
            for nm, fares, extra, fixed in rows:
                if fixed:
                    rec = {"origin": oid, "from": fixed, "to": nm}
                    if rec not in notes["normalizations"]:
                        notes["normalizations"].append(rec)
                # a hotel must live in ONE zone per origin
                for zr2, z2 in o["zones"].items():
                    if zr2 != zraw and nm in z2["hotels"]:
                        raise SystemExit(f"{oid}: {nm} appears in zones {zr2} and {zraw}")
                h = z["hotels"].get(nm)
                if h is None:
                    h = z["hotels"][nm] = {"fares": {}, "extra": {}, "order": len(z["hotels"])}
                    z["order"].append(nm)
                else:
                    # same hotel again on the SAME sheet (e.g. HARD ROCK twice in SDQ) or on the sister sheet
                    same_sheet = any(c in h["fares"] for c in fares)
                    if same_sheet:
                        assert all(h["fares"][c] == fares[c] for c in fares), f"{oid} {nm}: conflicting duplicate"
                        notes["duplicatesDropped"].append({"origin": oid, "hotel": nm, "sheet": sheet.strip()})
                        continue
                h["fares"].update(fares)
                h["extra"].update(extra)

    out_origins = []
    for oid in ("PUJ", "SDQ"):
        o = origins[oid]
        zones_out, total = [], 0
        for zraw in o["order"]:
            z = o["zones"][zraw]
            hotels = []
            for nm in z["order"]:
                h = z["hotels"][nm]
                hotels.append({
                    "id": slug(nm), "name": display_name(nm), "key": nm,
                    "fares": {c: int(h["fares"][c]) if float(h["fares"][c]).is_integer() else h["fares"][c]
                              for c in ("starex", "hiace", "suburban", "cadillac") if c in h["fares"]},
                    "extra": {c: int(h["extra"][c]) if float(h["extra"][c]).is_integer() else h["extra"][c]
                              for c in ("starex", "hiace", "suburban", "cadillac") if c in h["extra"]},
                })
            total += len(hotels)
            zones_out.append({"id": slug(zlabel(zraw)), "name": zlabel(zraw), "sheetName": zraw, "hotels": hotels})
        out_origins.append(dict(ORIGINS[oid], classes=["starex", "hiace", "suburban", "cadillac"],
                                zones=zones_out, hotelCount=total))
    # ---- sanity flags worth confirming with the client (data is kept verbatim) ----
    for o in out_origins:
        for z in o["zones"]:
            for h in z["hotels"]:
                ex = h["extra"]
                if len(set(ex.values())) > 1:
                    notes["flags"].append({"origin": o["id"], "hotel": h["name"],
                                           "note": "Pax adicional differs between sheets: " + json.dumps(ex)})
    return out_origins, notes


def pdf_rows(pdf):
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True, check=True).stdout
    rx = re.compile(r"^\s*(.+?)\s+\$\s*([\d.,]+)\s+\$\s*([\d.,]+)\s+\$\s*([\d.,]+)\s*$")
    rows = []
    for line in txt.splitlines():
        m = rx.match(line)
        if m:
            nm, _ = clean_name(m.group(1))
            rows.append((nm, float(m.group(2)), float(m.group(3)), float(m.group(4))))
    return rows


PDF_MAP = {  # substring of file name -> (origin, classes)
    "SUB_URBAN_CADILLAC": ("PUJ", ("suburban", "cadillac")),
    "MINIVAN__TECHO_ALTO__1_": ("PUJ", ("starex", "hiace")),
    "MINIVAN__TECHO_ALTO": ("PUJ", ("starex", "hiace")),
    "AERO_PUERTO_SANTO_DOMINGO_SUBURBAN": ("SDQ", ("suburban", "cadillac")),
    "AERO_PUERTO_SANTO_DOMINGO_MINIVAN": ("SDQ", ("starex", "hiace")),
}


def verify_pdfs(origins, pdf_dir):
    idx = {o["id"]: {h["key"]: h for z in o["zones"] for h in z["hotels"]} for o in origins}
    report = []
    for pdf in sorted(glob.glob(os.path.join(pdf_dir, "*.pdf"))):
        base = os.path.basename(pdf)
        key = next((k for k in ("MINIVAN__TECHO_ALTO__1_", "SUB_URBAN_CADILLAC", "MINIVAN__TECHO_ALTO",
                                "AERO_PUERTO_SANTO_DOMINGO_SUBURBAN", "AERO_PUERTO_SANTO_DOMINGO_MINIVAN") if k in base), None)
        if not key:
            continue
        oid, (c1, c2) = PDF_MAP[key]
        bad, n = [], 0
        seen = set()
        for nm, a, b, e in pdf_rows(pdf):
            n += 1
            h = idx[oid].get(nm)
            if not h:
                bad.append((nm, "not in JSON")); continue
            if (h["fares"].get(c1), h["fares"].get(c2), h["extra"].get(c1)) != (a, b, e) and \
               (float(h["fares"].get(c1, -1)), float(h["fares"].get(c2, -1)), float(h["extra"].get(c1, -1))) != (a, b, e):
                bad.append((nm, f"pdf {a}/{b}/{e} vs json {h['fares'].get(c1)}/{h['fares'].get(c2)}/{h['extra'].get(c1)}"))
            seen.add(nm)
        missing = [k for k, h in idx[oid].items() if c1 in h["fares"] and k not in seen]
        report.append({"pdf": base, "origin": oid, "classes": [c1, c2], "rows": n, "mismatches": bad, "jsonHotelsNotInPdf": missing})
    return report


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsm", required=True)
    ap.add_argument("--pdf-dir")
    ap.add_argument("--out-dir", default=os.path.join(os.path.dirname(__file__), "..", "assets"))
    a = ap.parse_args()
    origins, notes = build(a.xlsm)
    doc = {
        "title": "Tarifario de traslados regular",
        "currency": "USD",
        "updated": datetime.date.today().isoformat(),
        "unit": "Tarifa por vía (one-way), por vehículo, en USD",
        "sources": [
            "Tarifario_regular_Karun_Travel__Definitivo.xlsm (hojas DESDE PUNTA CANA SUB CADILLAC · DESDE AEROP PUNTA CANA MINIVAN · DESDE AEROP SDQ · DESDE AEROP SDQ SUB CAD)",
            "PDF: Suburban+Cadillac PUJ · Starex+Techo Alto PUJ · Starex Hyundai+Hiace Toyota PUJ · Suburban+Cadillac SDQ · Starex Hyundai+Hiace Toyota SDQ",
        ],
        "pricingModel": {
            "base": "La tarifa base cubre hasta paxMax del vehículo (1 vía).",
            "extraPax": "Cada pasajero adicional por vía se cobra el valor 'extra' de esa hoja (Starex/Hiace y Suburban/Cadillac son tablas distintas).",
            "multiLeg": "Demo: cada tramo adicional (regreso, otro traslado) se cotiza como una vía con la misma tarifa. No hay descuentos por ida y vuelta en las hojas.",
        },
        "vehicles": VEHICLES,
        "origins": origins,
        "meta": notes,
    }
    out = os.path.abspath(a.out_dir)
    with open(os.path.join(out, "tarifario.json"), "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
        f.write("\n")
    compact = json.dumps(doc, ensure_ascii=False, separators=(",", ":"))
    with open(os.path.join(out, "tarifario.data.js"), "w", encoding="utf-8") as f:
        f.write("/* GENERATED by tools/build_tarifario.py — do not edit. Source of truth: assets/tarifario.json */\n")
        f.write("window.KRN_TARIFARIO=" + compact + ";\n")
    for o in origins:
        print(o["id"], o["hotelCount"], "hoteles ·", len(o["zones"]), "zonas")
    print("normalizaciones:", notes["normalizations"])
    print("duplicados descartados:", notes["duplicatesDropped"])
    print("banderas:", json.dumps(notes["flags"], ensure_ascii=False))
    if a.pdf_dir:
        rep = verify_pdfs(origins, a.pdf_dir)
        ok = True
        for r in rep:
            print(f"PDF {r['pdf'][-60:]}: {r['rows']} filas · {len(r['mismatches'])} diferencias · sin PDF: {len(r['jsonHotelsNotInPdf'])}")
            for m in r["mismatches"][:20]:
                print("   ✗", m); ok = False
            for m in r["jsonHotelsNotInPdf"][:20]:
                print("   ? en JSON pero no en PDF:", m)
        print("VERIFICACIÓN PDF:", "OK" if ok else "CON DIFERENCIAS")
        if not ok:
            sys.exit(2)


if __name__ == "__main__":
    main()
